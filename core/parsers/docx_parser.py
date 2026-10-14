"""Word DOCX table extraction parser."""

import io
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional, Tuple, Union
import docx

from config.settings import SlaRulesConfig, get_sla_rules, resolve_column_name
from core.models.sla import ExtractionResult, SlaRecord
from core.parsers.html_parser import clean_cell_text, parse_date_value


def _find_docx_header_row(
    table: Any, mappings: Any
) -> Tuple[Optional[int], List[Tuple[str, Optional[str]]]]:
    """Scan the top rows of a DOCX table to identify the header row containing ticket_id."""
    max_scan = min(5, len(table.rows))
    for r_idx in range(max_scan):
        row = table.rows[r_idx]
        if not row.cells:
            continue

        candidate_headers: List[Tuple[str, Optional[str]]] = []
        has_ticket_id = False
        seen_headers: Dict[str, int] = {}

        for idx, cell in enumerate(row.cells):
            raw_title = clean_cell_text(cell.text)
            if not raw_title:
                col_name = f"Column_{idx + 1}"
            else:
                if raw_title in seen_headers:
                    seen_headers[raw_title] += 1
                    col_name = f"{raw_title}_{seen_headers[raw_title]}"
                else:
                    seen_headers[raw_title] = 1
                    col_name = raw_title

            canonical = resolve_column_name(raw_title, mappings)
            if canonical == "ticket_id":
                has_ticket_id = True
            candidate_headers.append((col_name, canonical))

        if has_ticket_id:
            return r_idx, candidate_headers

    # Fallback to row 0 if table has rows
    if len(table.rows) > 0:
        row = table.rows[0]
        candidate_headers = []
        seen_headers = {}
        for idx, cell in enumerate(row.cells):
            raw_title = clean_cell_text(cell.text)
            if not raw_title:
                col_name = f"Column_{idx + 1}"
            else:
                if raw_title in seen_headers:
                    seen_headers[raw_title] += 1
                    col_name = f"{raw_title}_{seen_headers[raw_title]}"
                else:
                    seen_headers[raw_title] = 1
                    col_name = raw_title
            canonical = resolve_column_name(raw_title, mappings)
            candidate_headers.append((col_name, canonical))
        return 0, candidate_headers

    return None, []


def parse_docx_table(
    docx_source: Union[Path, str, bytes, BinaryIO],
    source_filename: str = "upload.docx",
    rules: Optional[SlaRulesConfig] = None,
) -> ExtractionResult:
    """Extract ticket records from Word DOCX tables into structured SlaRecords.

    Args:
        docx_source: File path, byte buffer, or file-like stream of the DOCX document.
        source_filename: Name of the uploaded or source file.
        rules: Optional custom SLA rules configuration. Uses default if None.

    Returns:
        ExtractionResult containing parsed records, row counts, warnings, and errors.
    """
    rules = rules or get_sla_rules()
    warnings: List[str] = []
    errors: List[str] = []
    parsed_records: List[SlaRecord] = []
    unparsed_rows = 0

    if isinstance(docx_source, bytes):
        source_input = io.BytesIO(docx_source)
    elif isinstance(docx_source, Path):
        source_input = str(docx_source)
    else:
        source_input = docx_source

    try:
        doc = docx.Document(source_input)
    except Exception as exc:
        errors.append(f"Failed to open DOCX document: {exc}")
        return ExtractionResult(
            source_filename=source_filename,
            file_type="docx",
            total_rows=0,
            parsed_records=[],
            unparsed_rows=0,
            warnings=warnings,
            errors=errors,
        )

    if not doc.tables:
        warnings.append("No tables found in document")
        return ExtractionResult(
            source_filename=source_filename,
            file_type="docx",
            total_rows=0,
            parsed_records=[],
            unparsed_rows=0,
            warnings=warnings,
            errors=errors,
        )

    ticket_tables_found = 0

    for table_idx, table in enumerate(doc.tables):
        if not table.rows:
            continue

        header_idx, headers = _find_docx_header_row(table, rules.column_mappings)
        has_ticket_key = any(canon == "ticket_id" for _, canon in headers)

        if not has_ticket_key:
            # Skip tables without ticket ID column
            continue

        ticket_tables_found += 1
        data_rows = table.rows[header_idx + 1 :] if header_idx is not None else []

        for row_pos, row in enumerate(data_rows, start=header_idx + 2 if header_idx is not None else 1):
            cell_texts = [clean_cell_text(c.text) for c in row.cells]

            # Skip completely empty rows
            if not any(cell_texts):
                continue

            # Skip repeated header row
            if all(
                clean_cell_text(row.cells[i].text) == headers[i][0]
                for i in range(min(len(row.cells), len(headers)))
            ):
                continue

            row_data: Dict[str, Any] = {}
            raw_fields: Dict[str, Any] = {}

            for col_idx, (orig_header, canon_field) in enumerate(headers):
                cell_val = cell_texts[col_idx] if col_idx < len(cell_texts) else ""

                if canon_field == "ticket_id":
                    row_data["key"] = cell_val
                elif canon_field == "summary":
                    row_data["summary"] = cell_val or None
                elif canon_field == "team":
                    row_data["team"] = cell_val or None
                elif canon_field == "status":
                    row_data["status"] = cell_val or None
                elif canon_field == "priority":
                    row_data["priority"] = cell_val or None
                elif canon_field in ("created_date", "resolved_date", "due_date"):
                    dt = parse_date_value(
                        cell_val,
                        rules.date_parsing.formats,
                        rules.date_parsing.default_timezone,
                    )
                    row_data[canon_field] = dt
                    if dt is None and cell_val:
                        raw_fields[orig_header] = cell_val
                else:
                    # Unmapped column preserved in raw_fields
                    raw_fields[orig_header] = cell_val

            # Preserve extra cells beyond known header count
            if len(cell_texts) > len(headers):
                for extra_idx in range(len(headers), len(cell_texts)):
                    raw_fields[f"Column_{extra_idx + 1}"] = cell_texts[extra_idx]

            ticket_key = row_data.get("key", "").strip()
            if not ticket_key:
                unparsed_rows += 1
                warnings.append(
                    f"Table {table_idx + 1}, Row {row_pos}: Blank or missing ticket identifier, skipped."
                )
                continue

            record = SlaRecord(
                key=ticket_key,
                summary=row_data.get("summary"),
                team=row_data.get("team"),
                status=row_data.get("status"),
                priority=row_data.get("priority"),
                created_date=row_data.get("created_date"),
                resolved_date=row_data.get("resolved_date"),
                due_date=row_data.get("due_date"),
                raw_fields=raw_fields,
            )
            parsed_records.append(record)

    if ticket_tables_found == 0:
        warnings.append("No table containing a recognized ticket identifier column ('ticket_id') was found")

    total_rows = len(parsed_records) + unparsed_rows

    return ExtractionResult(
        source_filename=source_filename,
        file_type="docx",
        total_rows=total_rows,
        parsed_records=parsed_records,
        unparsed_rows=unparsed_rows,
        errors=errors,
        warnings=warnings,
    )


def parse_docx_file(
    file_path: Union[Path, str],
    rules: Optional[SlaRulesConfig] = None,
) -> ExtractionResult:
    """Read a DOCX file from disk and parse tables into ExtractionResult.

    Args:
        file_path: Path to the DOCX file.
        rules: Optional custom SLA rules configuration.

    Returns:
        ExtractionResult containing parsed records or file read errors.
    """
    path = Path(file_path)
    if not path.is_file():
        return ExtractionResult(
            source_filename=path.name,
            file_type="docx",
            total_rows=0,
            parsed_records=[],
            unparsed_rows=0,
            errors=[f"File not found: {path}"],
            warnings=[],
        )

    try:
        with open(path, "rb") as f:
            content = f.read()
    except Exception as exc:
        return ExtractionResult(
            source_filename=path.name,
            file_type="docx",
            total_rows=0,
            parsed_records=[],
            unparsed_rows=0,
            errors=[f"Failed to read file {path}: {exc}"],
            warnings=[],
        )

    return parse_docx_table(content, source_filename=path.name, rules=rules)
