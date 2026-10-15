"""Comprehensive unit tests covering edge cases across parsers, SLA engine, severity, analytics, and API."""

from datetime import datetime, timedelta, timezone
import io
import docx
from fastapi.testclient import TestClient
import pytest

from config.settings import AppSettings, SlaRulesConfig, get_settings
from core.analytics.metrics import (
    compute_analytics_payload,
    compute_overdue_distribution,
    compute_severity_distribution,
    compute_team_rankings,
)
from core.models.sla import SeverityLevel, SlaRecord
from core.parsers.date_parser import parse_date_safe, parse_datetime
from core.parsers.docx_parser import parse_docx_table
from core.parsers.html_parser import parse_html_table
from core.sla.engine import SlaEngine
from core.sla.severity import evaluate_severity
from web.app import create_app


# ==============================================================================
# 1. Date Parser Boundary & Edge Case Unit Tests
# ==============================================================================

def test_date_parser_leap_year() -> None:
    """Verify leap day boundary handling for 2024-02-29."""
    dt = parse_datetime("2024-02-29 23:59:59")
    assert dt is not None
    assert dt.year == 2024 and dt.month == 2 and dt.day == 29
    assert dt.hour == 23 and dt.minute == 59 and dt.second == 59
    assert dt.tzinfo == timezone.utc


def test_date_parser_microsecond_precision() -> None:
    """Verify ISO string with microsecond fraction is parsed into UTC datetime."""
    dt = parse_datetime("2026-07-20T15:30:45.987654Z")
    assert dt is not None
    assert dt.year == 2026 and dt.month == 7 and dt.day == 20
    assert dt.microsecond == 987654
    assert dt.tzinfo == timezone.utc


def test_date_parser_timezone_offsets() -> None:
    """Verify timezone offsets are accurately converted to UTC normalized time."""
    # +05:30 (India Standard Time) -> 12:30 IST is 07:00 UTC
    dt_ist = parse_datetime("2026-10-15T12:30:00+05:30")
    assert dt_ist is not None
    assert dt_ist.tzinfo == timezone.utc
    assert dt_ist.hour == 7 and dt_ist.minute == 0

    # -04:00 (EDT) -> 10:00 EDT is 14:00 UTC
    dt_edt = parse_datetime("2026-10-15T10:00:00-04:00")
    assert dt_edt is not None
    assert dt_edt.tzinfo == timezone.utc
    assert dt_edt.hour == 14 and dt_edt.minute == 0


def test_date_parser_empty_and_garbage_inputs() -> None:
    """Verify unparseable strings, N/A, and empty values return None safely without raising exceptions."""
    assert parse_datetime(None) is None
    assert parse_datetime("") is None
    assert parse_datetime("   ") is None
    assert parse_datetime("\t\n") is None
    assert parse_datetime("N/A") is None
    assert parse_datetime("TBD") is None
    assert parse_datetime("NULL") is None
    assert parse_datetime("99/99/9999") is None
    assert parse_datetime("invalid-date-format") is None


# ==============================================================================
# 2. SLA Engine & Severity Boundary Unit Tests
# ==============================================================================

def test_severity_boundary_conditions() -> None:
    """Verify exact boundary percentage transitions for severity evaluation."""
    rules = SlaRulesConfig().severity_rules

    # Exactly at warning threshold (80%), overdue = 0.0
    sev_80 = evaluate_severity(elapsed_days=8.0, target_days=10.0, overdue_days=0.0, rules=rules)
    assert sev_80 == SeverityLevel.WARNING

    # Just below warning threshold (79.9%) -> OK
    sev_79 = evaluate_severity(elapsed_days=7.99, target_days=10.0, overdue_days=0.0, rules=rules)
    assert sev_79 == SeverityLevel.OK

    # Overdue > 3.0 days -> CRITICAL
    sev_critical = evaluate_severity(elapsed_days=15.0, target_days=10.0, overdue_days=3.5, rules=rules)
    assert sev_critical == SeverityLevel.CRITICAL

    # Overdue > 0.0 days but <= 3.0 days -> BREACHED
    sev_breached = evaluate_severity(elapsed_days=11.5, target_days=10.0, overdue_days=1.5, rules=rules)
    assert sev_breached == SeverityLevel.BREACHED

    # Negative elapsed (resolved before created) -> OK
    sev_neg = evaluate_severity(elapsed_days=-1.5, target_days=5.0, overdue_days=0.0, rules=rules)
    assert sev_neg == SeverityLevel.OK



def test_sla_engine_resolved_before_created() -> None:
    """Verify SLA engine safely calculates tickets where resolved_date < created_date."""
    engine = SlaEngine()
    created = datetime(2026, 10, 10, 12, 0, tzinfo=timezone.utc)
    resolved = datetime(2026, 10, 10, 10, 0, tzinfo=timezone.utc)  # 2 hours earlier

    record = SlaRecord(
        key="TEST-NEG",
        priority="P1",
        created_date=created,
        resolved_date=resolved,
    )
    evaluated = engine.calculate_record(record)

    # Elapsed should be clamped or reflect <= 0 days, and should not be breached
    assert evaluated.is_breached is False
    assert evaluated.severity == SeverityLevel.OK


def test_sla_engine_extreme_durations() -> None:
    """Verify tickets created 10+ years ago calculate overdue days without numeric errors."""
    engine = SlaEngine()
    created = datetime(2014, 1, 1, 0, 0, tzinfo=timezone.utc)
    now = datetime(2026, 1, 1, 0, 0, tzinfo=timezone.utc)

    record = SlaRecord(
        key="ANCIENT-1",
        priority="P1",
        created_date=created,
    )
    evaluated = engine.calculate_record(record, as_of=now)

    assert evaluated.is_breached is True
    assert evaluated.severity == SeverityLevel.CRITICAL
    assert evaluated.elapsed_days > 4000.0
    assert evaluated.overdue_days > 4000.0


def test_sla_engine_summary_all_compliant() -> None:
    """Verify summary metrics when 100% of tickets are fully compliant."""
    engine = SlaEngine()
    created = datetime(2026, 10, 10, tzinfo=timezone.utc)
    records = [
        SlaRecord(key=f"OK-{i}", priority="P4", created_date=created, resolved_date=created + timedelta(hours=1))
        for i in range(5)
    ]
    batch = engine.calculate_batch(records)
    summary = engine.calculate_summary_metrics(batch)

    assert summary.total_records == 5
    assert summary.breached_count == 0
    assert summary.breach_rate_pct == 0.0
    assert summary.avg_overdue_days == 0.0
    assert summary.by_severity["OK"] == 5
    assert summary.by_severity["CRITICAL"] == 0


# ==============================================================================
# 3. HTML Parser Unicode, Emojis & Malformed Structure Tests
# ==============================================================================

def test_html_parser_unicode_and_emojis() -> None:
    """Verify parser extracts unicode summaries, emojis, and right-to-left scripts."""
    html = """
    <table>
        <tr>
            <th>Key</th>
            <th>Summary</th>
            <th>Priority</th>
            <th>Status</th>
            <th>Team</th>
        </tr>
        <tr>
            <td>SEC-999</td>
            <td>🔥🚨 Critical Database Crash in Primary Cluster!</td>
            <td>P1</td>
            <td>Open</td>
            <td>سائبر سیکیورٹی (CyberSec)</td>
        </tr>
        <tr>
            <td>INT-888</td>
            <td>Überprüfung der Netzwerkverbindung &amp; Latency</td>
            <td>P2</td>
            <td>In Progress</td>
            <td>Netzwerk Ops</td>
        </tr>
    </table>
    """
    res = parse_html_table(html, source_filename="unicode_test.html")
    assert len(res.parsed_records) == 2
    rec1 = res.parsed_records[0]
    assert rec1.key == "SEC-999"
    assert "🔥🚨" in rec1.summary
    assert "سائبر سیکیورٹی" in rec1.team

    rec2 = res.parsed_records[1]
    assert rec2.key == "INT-888"
    assert "Überprüfung" in rec2.summary


def test_html_parser_td_only_headers() -> None:
    """Verify HTML tables that use <td> in first row instead of <th> are parsed successfully."""
    html = """
    <table>
        <tr>
            <td>Issue Key</td>
            <td>Description</td>
            <td>Urgency</td>
        </tr>
        <tr>
            <td>LEGACY-1</td>
            <td>Legacy system migration</td>
            <td>High</td>
        </tr>
    </table>
    """
    res = parse_html_table(html)
    assert len(res.parsed_records) == 1
    assert res.parsed_records[0].key == "LEGACY-1"
    assert res.parsed_records[0].summary == "Legacy system migration"


# ==============================================================================
# 4. DOCX Parser Edge Cases
# ==============================================================================

def test_docx_parser_header_only_table() -> None:
    """Verify DOCX table with only header and 0 data rows returns 0 parsed records without errors."""
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Key"
    table.rows[0].cells[1].text = "Summary"
    table.rows[0].cells[2].text = "Priority"

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    res = parse_docx_table(stream)
    assert res.total_rows == 0
    assert len(res.parsed_records) == 0
    assert len(res.errors) == 0


def test_docx_parser_empty_rows_filtered() -> None:
    """Verify DOCX table with empty whitespace rows skips blank rows cleanly."""
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Key"
    table.rows[0].cells[1].text = "Summary"
    # Row 1 has valid ticket
    table.rows[1].cells[0].text = "DOC-VALID"
    table.rows[1].cells[1].text = "Valid Ticket Summary"
    # Row 2 is completely blank whitespace
    table.rows[2].cells[0].text = "   "
    table.rows[2].cells[1].text = "\t  \n"

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    res = parse_docx_table(stream)
    assert len(res.parsed_records) == 1
    assert res.parsed_records[0].key == "DOC-VALID"


# ==============================================================================
# 5. Analytics Distribution & Team Ranking Edge Cases
# ==============================================================================

def test_analytics_unassigned_team_grouping() -> None:
    """Verify tickets with None or whitespace team names are grouped as Unassigned."""
    records = [
        SlaRecord(key="T-1", team=None, is_breached=True, overdue_days=2.0),
        SlaRecord(key="T-2", team="   ", is_breached=False),
        SlaRecord(key="T-3", team="SecOps", is_breached=True, overdue_days=3.0),
    ]
    rankings = compute_team_rankings(records)

    team_names = [t["team"] for t in rankings]
    assert "Unassigned" in team_names
    assert "SecOps" in team_names

    unassigned = next(t for t in rankings if t["team"] == "Unassigned")
    assert unassigned["total"] == 2
    assert unassigned["breached"] == 1


def test_analytics_max_overdue_when_no_breaches() -> None:
    """Verify max_overdue_days is 0.0 when tickets exist with zero or negative overdue values."""
    records = [
        SlaRecord(key="T-1", overdue_days=None, is_breached=False),
        SlaRecord(key="T-2", overdue_days=0.0, is_breached=False),
    ]
    payload = compute_analytics_payload(records)
    assert payload["summary"]["max_overdue_days"] == 0.0
    assert payload["summary"]["breached_count"] == 0


# ==============================================================================
# 6. REST API File Ingestion Rejections & Safety
# ==============================================================================

@pytest.fixture
def client() -> TestClient:
    app = create_app()
    return TestClient(app)


@pytest.mark.parametrize("ext", [".csv", ".xlsx", ".json", ".zip", ".txt", ".exe"])
def test_api_upload_unsupported_extensions(client: TestClient, ext: str) -> None:
    """Verify that unsupported file extensions are rejected with 400 Bad Request."""
    dummy_file = io.BytesIO(b"sample file content")
    response = client.post(
        "/api/upload",
        files={"file": (f"test_file{ext}", dummy_file, "application/octet-stream")},
    )
    assert response.status_code == 400
    data = response.json()
    assert f"Unsupported file format '{ext}'" in data["detail"]
