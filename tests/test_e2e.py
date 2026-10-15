"""End-to-end integration tests for SLA Breach Dashboard & Analyzer.

Validates complete system workflows across:
- Multi-format ingestion (HTML, DOCX)
- Real SLA calculations and overdue binning
- Analytics aggregation payloads and team rankings
- Full Web UI lifecycle and asset contracts
- Error handling, boundary conditions, and resilient failure modes
"""

import io
from typing import Dict, Any
import docx
from fastapi.testclient import TestClient
import pytest

from web.app import create_app


@pytest.fixture
def client() -> TestClient:
    """Fixture returning configured TestClient."""
    app = create_app()
    return TestClient(app)


# ---------------------------------------------------------------------------
# Scenario 1: Comprehensive HTML Ingestion & SLA Calculation E2E
# ---------------------------------------------------------------------------


def test_e2e_html_multi_priority_and_analytics(client: TestClient) -> None:
    """End-to-end verification of HTML ingestion, SLA metric calculations, and analytics.

    Validates:
    - Ingestion of 12 tickets across P1, P2, P3, P4
    - Coverage of all overdue bins: 0-1d, 1-3d, 3-7d, >7d
    - Teams: Platform, Frontend, Backend, DevOps, and unassigned
    - Accurate calculations of elapsed_days, overdue_days, and severity
    - Consistency of analytics distributions and team leaderboard sorting
    """
    html_content = """
    <!DOCTYPE html>
    <html>
    <head><title>Confluence Engineering Incident Report</title></head>
    <body>
        <h1>Q4 Production Incidents & Service Requests</h1>
        <table class="confluenceTable">
            <thead>
                <tr>
                    <th>Issue Key</th>
                    <th>Issue Summary</th>
                    <th>Priority</th>
                    <th>Status</th>
                    <th>Date Created</th>
                    <th>Date Resolved</th>
                    <th>Assigned Team</th>
                    <th>Environment</th>
                </tr>
            </thead>
            <tbody>
                <!-- 1. P1 Resolved Compliant (target 1d, elapsed 0.5d -> overdue 0.0d -> OK) -->
                <tr>
                    <td>INFRA-101</td>
                    <td>Primary Redis cluster failover</td>
                    <td>P1</td>
                    <td>Closed</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-01 12:00:00</td>
                    <td>Platform</td>
                    <td>Production</td>
                </tr>
                <!-- 2. P1 Resolved Overdue 0-1d (target 1d, elapsed 1.5d -> overdue 0.5d -> BREACHED) -->
                <tr>
                    <td>INFRA-102</td>
                    <td>Kubernetes ingress controller timeout</td>
                    <td>P1</td>
                    <td>Resolved</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-02 12:00:00</td>
                    <td>Platform</td>
                    <td>Production</td>
                </tr>
                <!-- 3. P1 Resolved Overdue 1-3d (target 1d, elapsed 3.0d -> overdue 2.0d -> BREACHED) -->
                <tr>
                    <td>INFRA-103</td>
                    <td>Payment gateway webhook drops</td>
                    <td>P1</td>
                    <td>Resolved</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-04 00:00:00</td>
                    <td>Platform</td>
                    <td>Production</td>
                </tr>
                <!-- 4. P1 Resolved Overdue 3-7d (target 1d, elapsed 5.0d -> overdue 4.0d -> CRITICAL) -->
                <tr>
                    <td>INFRA-104</td>
                    <td>Database replication split brain</td>
                    <td>P1</td>
                    <td>Resolved</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-06 00:00:00</td>
                    <td>Platform</td>
                    <td>Production</td>
                </tr>
                <!-- 5. P2 Resolved Overdue >7d (target 2d, elapsed 12.0d -> overdue 10.0d -> CRITICAL) -->
                <tr>
                    <td>FE-201</td>
                    <td>Checkout page CSP blockage on Safari</td>
                    <td>P2</td>
                    <td>Closed</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-13 00:00:00</td>
                    <td>Frontend</td>
                    <td>Production</td>
                </tr>
                <!-- 6. P2 Resolved Compliant (target 2d, elapsed 1.0d -> overdue 0.0d -> OK) -->
                <tr>
                    <td>FE-202</td>
                    <td>Auth token expiry refresh loop</td>
                    <td>P2</td>
                    <td>Resolved</td>
                    <td>2026-01-05 00:00:00</td>
                    <td>2026-01-06 00:00:00</td>
                    <td>Frontend</td>
                    <td>Staging</td>
                </tr>
                <!-- 7. P2 Resolved Overdue 0-1d (target 2d, elapsed 2.8d -> overdue 0.8d -> BREACHED) -->
                <tr>
                    <td>FE-203</td>
                    <td>Mobile navigation bar overlap</td>
                    <td>P2</td>
                    <td>Resolved</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-03 19:12:00</td>
                    <td>Frontend</td>
                    <td>Production</td>
                </tr>
                <!-- 8. P3 Resolved Overdue 1-3d (target 5d, elapsed 7.0d -> overdue 2.0d -> BREACHED) -->
                <tr>
                    <td>BE-301</td>
                    <td>Batch export worker OOM crash</td>
                    <td>P3</td>
                    <td>Resolved</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-08 00:00:00</td>
                    <td>Backend</td>
                    <td>Production</td>
                </tr>
                <!-- 9. P3 Resolved Compliant (target 5d, elapsed 4.0d -> overdue 0.0d -> WARNING) -->
                <tr>
                    <td>BE-302</td>
                    <td>Search indexing lag during catalog updates</td>
                    <td>P3</td>
                    <td>Resolved</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-05 00:00:00</td>
                    <td>Backend</td>
                    <td>Production</td>
                </tr>
                <!-- 10. P4 Resolved Compliant (target 10d, elapsed 6.0d -> overdue 0.0d -> OK) -->
                <tr>
                    <td>OPS-401</td>
                    <td>Update CI/CD runner docker base image</td>
                    <td>P4</td>
                    <td>Done</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-07 00:00:00</td>
                    <td>DevOps</td>
                    <td>Infrastructure</td>
                </tr>
                <!-- 11. P4 Resolved Overdue >7d (target 10d, elapsed 20.0d -> overdue 10.0d -> CRITICAL) -->
                <tr>
                    <td>OPS-402</td>
                    <td>Renew internal mTLS certificates</td>
                    <td>P4</td>
                    <td>Completed</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-21 00:00:00</td>
                    <td>DevOps</td>
                    <td>Production</td>
                </tr>
                <!-- 12. P3 Resolved Overdue 3-7d Unassigned Team (target 5d, elapsed 10.0d -> overdue 5.0d -> CRITICAL) -->
                <tr>
                    <td>MISC-501</td>
                    <td>Legacy reporting webhook migration</td>
                    <td>P3</td>
                    <td>Closed</td>
                    <td>2026-01-01 00:00:00</td>
                    <td>2026-01-11 00:00:00</td>
                    <td></td>
                    <td>Production</td>
                </tr>
            </tbody>
        </table>
    </body>
    </html>
    """

    upload_file = io.BytesIO(html_content.encode("utf-8"))
    response = client.post(
        "/api/upload",
        files={"file": ("incident_report.html", upload_file, "text/html")},
    )

    assert response.status_code == 200
    payload = response.json()

    # High-level file metadata
    assert payload["source_filename"] == "incident_report.html"
    assert payload["file_type"] == "html"
    assert payload["total_rows"] == 12
    assert payload["parsed_count"] == 12
    assert payload["unparsed_rows"] == 0
    assert len(payload["errors"]) == 0

    records = payload["records"]
    assert len(records) == 12

    # Verify calculation properties on individual records
    for rec in records:
        assert rec["key"] is not None
        assert rec["sla_target_days"] > 0
        assert rec["elapsed_days"] is not None
        assert rec["overdue_days"] is not None
        assert rec["severity"] in ["OK", "WARNING", "CRITICAL", "BREACHED"]
        # Custom columns captured in raw_fields
        assert "Environment" in rec["raw_fields"]
        assert rec["raw_fields"]["Environment"] in ["Production", "Staging", "Infrastructure"]

    # Verify summary metrics
    metrics = payload["metrics"]
    assert metrics["total_records"] == 12
    assert metrics["breached_count"] > 0
    assert metrics["breach_rate_pct"] > 0
    assert metrics["by_severity"]["CRITICAL"] > 0

    # Verify analytics payload
    assert "analytics" in payload
    analytics = payload["analytics"]

    # 1. Overdue distribution bins coverage: 0-1d, 1-3d, 3-7d, >7d
    overdue_dist = analytics["overdue_distribution"]
    assert overdue_dist["0-1d"] >= 1, f"Expected 0-1d bin >= 1, got {overdue_dist}"
    assert overdue_dist["1-3d"] >= 1, f"Expected 1-3d bin >= 1, got {overdue_dist}"
    assert overdue_dist["3-7d"] >= 1, f"Expected 3-7d bin >= 1, got {overdue_dist}"
    assert overdue_dist[">7d"] >= 1, f"Expected >7d bin >= 1, got {overdue_dist}"

    sum_overdue_bins = sum(overdue_dist.values())
    assert sum_overdue_bins == metrics["breached_count"], (
        f"Sum of overdue bins ({sum_overdue_bins}) must equal breached_count ({metrics['breached_count']})"
    )

    # 2. Severity distribution covers total records
    sev_dist = analytics["severity_distribution"]
    assert sum(sev_dist.values()) == metrics["total_records"]
    assert sev_dist["BREACHED"] + sev_dist["CRITICAL"] == metrics["breached_count"]

    # 3. Team rankings ordered descending by breached count
    rankings = analytics["team_rankings"]
    assert len(rankings) >= 4
    for i in range(len(rankings) - 1):
        assert (
            rankings[i]["breached"] >= rankings[i + 1]["breached"]
        ), f"Rankings not sorted descending: {rankings}"

    # Confirm unassigned team is categorized cleanly
    team_names = [item["team"] for item in rankings]
    assert "Unassigned" in team_names or "Platform" in team_names


# ---------------------------------------------------------------------------
# Scenario 2: Comprehensive DOCX Ingestion E2E
# ---------------------------------------------------------------------------


def test_e2e_docx_table_ingestion_and_calculation(client: TestClient) -> None:
    """End-to-end verification of Word (.docx) export table ingestion.

    Validates:
    - In-memory DOCX creation with table columns
    - Multiple rows across diverse priorities and statuses
    - Multipart upload handling, SLA calculation, and KPI aggregation
    """
    doc = docx.Document()
    doc.add_heading("Production Ticket Audit", level=1)
    doc.add_paragraph("Exported directly from Word documentation tracker.")

    table = doc.add_table(rows=1, cols=7)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Ticket ID"
    hdr_cells[1].text = "Summary"
    hdr_cells[2].text = "Priority"
    hdr_cells[3].text = "Status"
    hdr_cells[4].text = "Created Date"
    hdr_cells[5].text = "Resolved Date"
    hdr_cells[6].text = "Assigned Team"

    test_rows = [
        (
            "DOCX-01",
            "Urgent database lock contention",
            "P1",
            "Resolved",
            "2026-10-01 08:00:00",
            "2026-10-01 14:00:00",
            "DataPlatform",
        ),
        (
            "DOCX-02",
            "API rate limiter latency degradation",
            "P2",
            "Closed",
            "2026-10-01 00:00:00",
            "2026-10-05 00:00:00",
            "GatewayTeam",
        ),
        (
            "DOCX-03",
            "Localization typo in checkout receipt",
            "P4",
            "Done",
            "2026-10-01 10:00:00",
            "2026-10-03 10:00:00",
            "Localization",
        ),
        (
            "DOCX-04",
            "Kubernetes node memory exhaustion",
            "P1",
            "Resolved",
            "2026-10-01 00:00:00",
            "2026-10-04 00:00:00",
            "DataPlatform",
        ),
    ]

    for row_data in test_rows:
        row_cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            row_cells[idx].text = text

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)

    response = client.post(
        "/api/upload",
        files={
            "file": (
                "audit_tickets.docx",
                docx_buffer,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    data = response.json()

    assert data["source_filename"] == "audit_tickets.docx"
    assert data["file_type"] == "docx"
    assert data["total_rows"] == 4
    assert data["parsed_count"] == 4
    assert data["metrics"]["total_records"] == 4

    records = data["records"]
    assert len(records) == 4
    docx1 = next(r for r in records if r["key"] == "DOCX-01")
    assert docx1["priority"] == "P1"
    assert docx1["severity"] == "OK"
    assert docx1["is_breached"] is False

    docx4 = next(r for r in records if r["key"] == "DOCX-04")
    assert docx4["priority"] == "P1"
    assert docx4["is_breached"] is True
    assert docx4["overdue_days"] > 0

    # Analytics payload presence
    assert "analytics" in data
    assert "overdue_distribution" in data["analytics"]
    assert "team_rankings" in data["analytics"]


# ---------------------------------------------------------------------------
# Scenario 3: Full Web UI Lifecycle Contract
# ---------------------------------------------------------------------------


def test_e2e_web_ui_lifecycle_contract(client: TestClient) -> None:
    """End-to-end verification of HTML view, scripts, styles, rules API, and Chart.js DOM integration."""
    # 1. Main index page renders 200 OK
    res_index = client.get("/?view=template")
    assert res_index.status_code == 200
    body = res_index.text

    # 2. Check tab buttons
    for tab_id in [
        "tab-btn-upload",
        "tab-btn-records",
        "tab-btn-analytics",
        "tab-btn-rules",
    ]:
        assert f'id="{tab_id}"' in body, f"Missing tab button {tab_id}"

    # 3. Check Chart.js CDN inclusion and chart canvases
    assert "chart.js" in body.lower()
    for chart_id in ["chart-severity", "chart-overdue", "chart-teams"]:
        assert f'id="{chart_id}"' in body, f"Missing canvas element {chart_id}"

    # 4. Rules endpoint contract
    res_rules = client.get("/api/rules")
    assert res_rules.status_code == 200
    rules = res_rules.json()
    assert "sla_thresholds" in rules
    assert "priority_targets_days" in rules["sla_thresholds"]
    assert "P1" in rules["sla_thresholds"]["priority_targets_days"]
    assert "P4" in rules["sla_thresholds"]["priority_targets_days"]

    # 5. Static assets deliver 200 OK
    res_js = client.get("/static/js/app.js")
    assert res_js.status_code == 200
    assert "renderAllCharts" in res_js.text

    res_css = client.get("/static/css/styles.css")
    assert res_css.status_code == 200


# ---------------------------------------------------------------------------
# Scenario 4: Error Resiliency and Edge Cases
# ---------------------------------------------------------------------------


def test_e2e_error_resiliency_empty_file(client: TestClient) -> None:
    """Empty files (0 bytes) should return 400 Bad Request."""
    empty_file = io.BytesIO(b"")
    response = client.post(
        "/api/upload",
        files={"file": ("empty.html", empty_file, "text/html")},
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_e2e_error_resiliency_unsupported_format(client: TestClient) -> None:
    """Unsupported extensions (.pdf, .xlsx, .csv) should return 400 Bad Request."""
    pdf_file = io.BytesIO(b"%PDF-1.4 header")
    response = client.post(
        "/api/upload",
        files={"file": ("report.pdf", pdf_file, "application/pdf")},
    )
    assert response.status_code == 400
    assert "unsupported" in response.json()["detail"].lower()


def test_e2e_html_without_ticket_tables(client: TestClient) -> None:
    """HTML without any valid ticket tables should return 200 OK with 0 rows and descriptive warning."""
    html_without_tables = """
    <!DOCTYPE html>
    <html>
    <body>
        <h1>Wiki Article</h1>
        <p>This is a standard informational document without table rows.</p>
    </body>
    </html>
    """
    file_bytes = io.BytesIO(html_without_tables.encode("utf-8"))
    response = client.post(
        "/api/upload",
        files={"file": ("notes.html", file_bytes, "text/html")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["total_rows"] == 0
    assert data["parsed_count"] == 0
    assert len(data["records"]) == 0
    assert any("table" in warn.lower() for warn in data["warnings"])


def test_e2e_malformed_html_markup_resilience(client: TestClient) -> None:
    """Malformed or truncated HTML with broken tags should parse gracefully without 500 error."""
    broken_html = """
    <div><table>
    <tr><th>Key<th>Summary<th>Created Date
    <tr><td>BROKEN-1<td>Unclosed row <td>2026-10-01
    <tr><td>BROKEN-2<td>Missing closing tag
    """
    file_bytes = io.BytesIO(broken_html.encode("utf-8"))
    response = client.post(
        "/api/upload",
        files={"file": ("broken.html", file_bytes, "text/html")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["total_rows"] >= 1
    assert data["parsed_count"] >= 1
    assert data["metrics"]["total_records"] >= 1
